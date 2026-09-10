#!/usr/bin/env python3
"""Insert live app screenshots into PPTX placeholder slides 6-10."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import copy

PPTX_PATH = "TrustMeRecycle_Pitch_Deck.pptx"

# Theme colors
DARK_BG   = RGBColor(0x0D, 0x0F, 0x1A)
GREEN     = RGBColor(0x2C, 0x5F, 0x2D)
LIME      = RGBColor(0x97, 0xBC, 0x62)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAY      = RGBColor(0xCC, 0xCC, 0xCC)

SW = Inches(13.33)
SH = Inches(7.5)

def clear_slide(slide):
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        sp_tree.remove(shape._element)

def add_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG

def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_label(slide, text, left, top, width, height,
              size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txb

def add_image(slide, path, left, top, width, height):
    return slide.shapes.add_picture(path, left, top, width, height)

def build_admin_slide(slide, title, subtitle, img_path, caption):
    """Landscape admin screenshot slide layout."""
    clear_slide(slide)
    add_bg(slide)

    # Thin green top accent bar
    add_rect(slide, 0, 0, SW, Inches(0.05), GREEN)

    # Title
    add_label(slide, title,
              left=Inches(0.4), top=Inches(0.1),
              width=Inches(9), height=Inches(0.7),
              size=28, bold=True, color=LIME)

    # Subtitle
    add_label(slide, subtitle,
              left=Inches(0.4), top=Inches(0.75),
              width=Inches(12.5), height=Inches(0.4),
              size=13, color=GRAY)

    # Screenshot — admin images are ~1.59:1 landscape
    # At height 5.6" → width = 5.6 * 1.59 = 8.9"
    img_h = Inches(5.6)
    img_w = Inches(8.9)
    img_left = (SW - img_w) / 2
    img_top = Inches(1.25)
    add_image(slide, img_path, img_left, img_top, img_w, img_h)

    # Caption bar at bottom
    add_rect(slide, 0, Inches(7.1), SW, Inches(0.4), RGBColor(0x1A, 0x28, 0x1A))
    add_label(slide, caption,
              left=Inches(0.4), top=Inches(7.1),
              width=Inches(12.5), height=Inches(0.4),
              size=11, color=GRAY, align=PP_ALIGN.LEFT)

def build_vendor_slide(slide, title, img_path1, cap1, img_path2, cap2):
    """Portrait mobile screenshot slide — two phones side by side."""
    clear_slide(slide)
    add_bg(slide)

    # Thin green top accent bar
    add_rect(slide, 0, 0, SW, Inches(0.05), GREEN)

    # Title
    add_label(slide, title,
              left=Inches(0.4), top=Inches(0.1),
              width=Inches(12.5), height=Inches(0.7),
              size=28, bold=True, color=LIME)

    # Two portrait screenshots — 1080×2400 → ratio 0.45
    img_h = Inches(5.9)
    img_w = Inches(5.9 * 1080 / 2400)   # ≈ 2.655"
    gap = Inches(0.8)
    total = img_w * 2 + gap
    left1 = (SW - total) / 2
    left2 = left1 + img_w + gap
    img_top = Inches(0.95)

    add_image(slide, img_path1, left1, img_top, img_w, img_h)
    add_image(slide, img_path2, left2, img_top, img_w, img_h)

    # Labels under each image
    label_top = img_top + img_h + Inches(0.08)
    add_label(slide, cap1,
              left=left1, top=label_top,
              width=img_w + Inches(0.3), height=Inches(0.45),
              size=11, color=LIME, align=PP_ALIGN.CENTER)
    add_label(slide, cap2,
              left=left2, top=label_top,
              width=img_w + Inches(0.3), height=Inches(0.45),
              size=11, color=LIME, align=PP_ALIGN.CENTER)

    # Side annotation (right of images)
    right_x = left2 + img_w + Inches(0.3)
    right_w = SW - right_x - Inches(0.2)
    if right_w > Inches(1.5):
        bullets = [
            "✓ Role-based login",
            "✓ Customer account grid",
            "✓ State › City › Site picker",
            "✓ DONE / PENDING status",
            "✓ Mark as Picked Up",
        ]
        add_label(slide, "\n".join(bullets),
                  left=right_x, top=Inches(2.0),
                  width=right_w, height=Inches(3.0),
                  size=12, color=WHITE)

def main():
    prs = Presentation(PPTX_PATH)

    slides = prs.slides

    # ── Slide 6: Admin — Pickup Requests (Customer → Admin workflow) ──
    build_admin_slide(
        slides[5],
        title="Admin Panel — Customer Pickup Requests",
        subtitle="Admin receives & manages scrap pickup requests from Blinkit, Zepto, Zomato, Amazon and Flipkart warehouses",
        img_path="/tmp/admin_pickup_requests.jpg",
        caption="Live data  ·  4 active requests: Completed | Truck Dispatched | Assigned | Pending Review  ·  Admin assigns to vendor site in one click",
    )

    # ── Slide 7: Admin — Collection Jobs (Dispatch & tracking) ──
    build_admin_slide(
        slides[6],
        title="Admin Panel — Collection Job Tracking",
        subtitle="Dispatch trucks, assign drivers, track scrap weight collected — full end-to-end job lifecycle",
        img_path="/tmp/admin_collection_jobs.jpg",
        caption="Live data  ·  South Godown: Truck Dispatched  ·  North Godown: Completed 4.80 MT  ·  2× Pending pickup  ·  'Not Picked Up' alert triggers follow-up",
    )

    # ── Slide 8: Admin — Sites Capacity (16 godowns, real-time alerts) ──
    build_admin_slide(
        slides[7],
        title="Admin Panel — Real-Time Site Capacity Monitoring",
        subtitle="16 active godown sites across Delhi, Bengaluru & Lucknow — automated overflow alerts when capacity exceeds 100%",
        img_path="/tmp/admin_sites.jpg",
        caption="Live data  ·  North Godown: 121.0% OVER CAPACITY (alert)  ·  South Godown: 52.8%  ·  Per-site View Report & Dispatch Truck actions",
    )

    # ── Slide 9: Vendor Mobile — Login + Collection Dashboard ──
    build_vendor_slide(
        slides[8],
        title="Vendor Mobile App — Role Select & Collection Dashboard",
        img_path1="/tmp/vendor_screen1.png",
        cap1="Role Select Screen",
        img_path2="/tmp/vendor_dashboard.png",
        cap2="Collection Dashboard",
    )

    # ── Slide 10: Vendor Mobile — Jobs List & Pickup Confirmation ──
    build_vendor_slide(
        slides[9],
        title="Vendor Mobile App — Assigned Pickups & Job Completion",
        img_path1="/tmp/vendor_jobs_list.png",
        cap1="Assigned Jobs List",
        img_path2="/tmp/vendor_active_job.png",
        cap2="Mark as Picked Up",
    )

    prs.save(PPTX_PATH)
    print("✓ Saved:", PPTX_PATH)

if __name__ == "__main__":
    main()
